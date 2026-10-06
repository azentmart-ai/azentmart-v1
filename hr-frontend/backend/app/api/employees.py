import csv
import io
from datetime import date
from openpyxl import load_workbook
from fastapi import APIRouter, Depends, HTTPException, UploadFile, File
from sqlalchemy.orm import Session
from ..database import get_db
from ..dependencies import current_user, require_roles
from ..models import Attendance, Document, Employee, Leave, Onboarding, OnboardingTask, Payroll, User
from ..schemas import EmployeeIn

router = APIRouter(prefix="/employees", tags=["Employees"])

DEFAULT_TASKS = [
    ("Candidate & offer", "Offer and joining details", "Pre-joining"),
    ("Employee record", "Create and verify the employee master record", "HR"),
    ("Documents", "Collect required employee documents", "Documents"),
    ("Background verification", "Complete background verification", "Verification"),
    ("IT access", "Provision required systems and access", "IT"),
    ("HR orientation", "Complete HR policies and orientation", "Orientation"),
    ("Manager & training", "Assign manager and initial training", "Training"),
    ("Completion", "Close onboarding and confirm completion", "Completion"),
]

def serialize_employee(e):
    return {
        "id": e.id,
        "employee_id": f"EMP-{e.id:05d}",
        "name": e.name,
        "email": e.email,
        "phone": e.phone,
        "department": e.department,
        "designation": e.designation,
        "location": e.location,
        "manager": e.manager,
        "employment_type": e.employment_type,
        "salary": e.salary,
        "date_of_birth": str(e.date_of_birth) if e.date_of_birth else None,
        "bank_name": e.bank_name,
        "bank_account": e.bank_account,
        "ifsc_code": e.ifsc_code,
        "status": e.status,
        "join_date": str(e.join_date) if e.join_date else None,
        "user_id": e.user_id,
    }

def ensure_onboarding(db, employee_id):
    onboarding = db.query(Onboarding).filter(Onboarding.employee_id == employee_id).first()
    if onboarding:
        return onboarding
    onboarding = Onboarding(employee_id=employee_id, progress=0, status="Not Started")
    db.add(onboarding)
    db.flush()
    for title, description, category in DEFAULT_TASKS:
        db.add(OnboardingTask(
            onboarding_id=onboarding.id,
            title=title,
            description=description,
            category=category,
            completed=False,
        ))
    return onboarding

@router.get("")
def list_employees(
    q: str | None = None,
    department: str | None = None,
    status: str | None = None,
    page: int = 1,
    page_size: int = 50,
    db: Session = Depends(get_db),
    user=Depends(current_user),
):
    query = db.query(Employee)
    if q:
        term = f"%{q}%"
        query = query.filter(
            (Employee.name.ilike(term))
            | (Employee.email.ilike(term))
            | (Employee.department.ilike(term))
            | (Employee.designation.ilike(term))
            | (Employee.location.ilike(term))
        )
    if department:
        query = query.filter(Employee.department == department)
    if status:
        query = query.filter(Employee.status == status)

    total = query.count()
    page = max(1, page)
    page_size = min(max(1, page_size), 200)
    rows = (
        query.order_by(Employee.id.desc())
        .offset((page - 1) * page_size)
        .limit(page_size)
        .all()
    )
    return {
        "items": [serialize_employee(e) for e in rows],
        "total": total,
        "page": page,
        "page_size": page_size,
    }

@router.post("")
def create_employee(
    payload: EmployeeIn,
    db: Session = Depends(get_db),
    user=Depends(require_roles("hr_admin", "hr_manager", "admin")),
):
    if db.query(Employee).filter(Employee.email == payload.email).first():
        raise HTTPException(409, "Employee email already exists")

    data = payload.model_dump()
    data["join_date"] = data.get("join_date") or date.today()

    linked_user = db.query(User).filter(User.email == payload.email).first()
    employee = Employee(**data, user_id=linked_user.id if linked_user else None)
    db.add(employee)
    db.flush()
    ensure_onboarding(db, employee.id)
    db.commit()
    db.refresh(employee)
    return serialize_employee(employee)

@router.patch("/{employee_id}")
def update_employee(
    employee_id: int,
    payload: EmployeeIn,
    db: Session = Depends(get_db),
    user=Depends(require_roles("hr_admin", "hr_manager", "admin")),
):
    employee = db.get(Employee, employee_id)
    if not employee:
        raise HTTPException(404, "Employee not found")

    duplicate = (
        db.query(Employee)
        .filter(Employee.email == payload.email, Employee.id != employee_id)
        .first()
    )
    if duplicate:
        raise HTTPException(409, "Another employee already uses this email")

    for key, value in payload.model_dump().items():
        if key == "join_date" and value is None:
            continue
        setattr(employee, key, value)
    linked_user = db.query(User).filter(User.email == payload.email).first()
    if linked_user:
        employee.user_id = linked_user.id

    db.commit()
    db.refresh(employee)
    return serialize_employee(employee)

@router.delete("/{employee_id}")
def delete_employee(
    employee_id: int,
    db: Session = Depends(get_db),
    user=Depends(require_roles("hr_admin", "admin")),
):
    employee = db.get(Employee, employee_id)
    if not employee:
        raise HTTPException(404, "Employee not found")
    db.delete(employee)
    db.commit()
    return {"message": "Employee deleted"}

@router.post("/bulk")
def bulk_import(
    file: UploadFile = File(...),
    db: Session = Depends(get_db),
    user=Depends(require_roles("hr_admin", "hr_manager", "admin")),
):
    filename = (file.filename or "").lower()
    raw = file.file.read()

    if filename.endswith(".csv"):
        try:
            text = raw.decode("utf-8-sig")
        except Exception as exc:
            raise HTTPException(400, "Unable to read the CSV file.") from exc
        reader = csv.DictReader(io.StringIO(text))
    elif filename.endswith(".xlsx"):
        try:
            wb = load_workbook(io.BytesIO(raw), read_only=True, data_only=True)
            ws = wb.active
            rows = list(ws.iter_rows(values_only=True))
            headers_row = [str(x or "").strip().lower() for x in (rows[0] if rows else [])]
            records = [
                dict(zip(headers_row, [str(v or "") for v in row]))
                for row in rows[1:]
            ]
            class ExcelReader:
                fieldnames = headers_row
                def __iter__(self):
                    return iter(records)
            reader = ExcelReader()
        except Exception as exc:
            raise HTTPException(400, "Unable to read the Excel file.") from exc
    else:
        raise HTTPException(400, "Bulk import accepts .csv or .xlsx files.")

    required = {"name", "email"}
    headers = {str(x or "").strip().lower() for x in (reader.fieldnames or [])}
    if not required.issubset(headers):
        raise HTTPException(400, "CSV/XLSX must contain at least name and email columns.")

    created = []
    errors = []
    for n, row in enumerate(reader, start=2):
        clean = {
            str(k).strip().lower(): (str(v).strip() if v is not None else "")
            for k, v in row.items()
        }
        if not clean.get("name") or not clean.get("email"):
            errors.append({"row": n, "error": "Name and email are required."})
            continue
        if db.query(Employee).filter(Employee.email == clean["email"]).first():
            errors.append({"row": n, "error": f"Employee email already exists: {clean['email']}"})
            continue
        try:
            linked_user = db.query(User).filter(User.email == clean["email"]).first()
            employee = Employee(
                user_id=linked_user.id if linked_user else None,
                name=clean["name"],
                email=clean["email"],
                phone=clean.get("phone") or None,
                department=clean.get("department") or None,
                designation=clean.get("designation") or clean.get("role") or None,
                location=clean.get("location") or None,
                manager=clean.get("manager") or None,
                employment_type=clean.get("employment_type") or "Full-time",
                salary=float(clean["salary"]) if clean.get("salary") else None,
                join_date=date.fromisoformat(clean["join_date"]) if clean.get("join_date") else date.today(),
                status=clean.get("status") or "Active",
            )
            db.add(employee)
            db.flush()
            ensure_onboarding(db, employee.id)
            created.append(employee.id)
        except Exception as exc:
            db.rollback()
            errors.append({"row": n, "error": str(exc)})
    db.commit()
    return {"created": len(created), "errors": errors}

@router.get("/{employee_id}")
def get_employee(
    employee_id: int,
    db: Session = Depends(get_db),
    user=Depends(current_user),
):
    e = db.get(Employee, employee_id)
    if not e:
        raise HTTPException(404, "Employee not found")

    result = serialize_employee(e)
    result["attendance"] = [
        {"date": str(x.work_date), "check_in": x.check_in, "check_out": x.check_out, "status": x.status}
        for x in db.query(Attendance).filter(Attendance.employee_id == e.id).order_by(Attendance.id.desc()).limit(50)
    ]
    result["leave"] = [
        {"leave_type": x.leave_type, "start_date": str(x.start_date), "end_date": str(x.end_date), "status": x.status, "reason": x.reason}
        for x in db.query(Leave).filter(Leave.employee_id == e.id).order_by(Leave.id.desc()).limit(50)
    ]
    result["payroll"] = [
        {"month": x.month, "net_pay": x.net_pay, "status": x.status}
        for x in db.query(Payroll).filter(Payroll.employee_id == e.id).order_by(Payroll.id.desc()).limit(50)
    ]
    result["documents"] = [
        {"id": x.id, "title": x.title, "status": x.status, "category": x.category}
        for x in db.query(Document).filter((Document.employee_id == e.id) | (Document.employee_id == None)).order_by(Document.id.desc()).limit(50)
    ]
    onboarding = db.query(Onboarding).filter(Onboarding.employee_id == e.id).first()
    result["onboarding"] = (
        {"id": onboarding.id, "status": onboarding.status, "progress": onboarding.progress}
        if onboarding else None
    )
    return result
